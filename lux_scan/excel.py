"""Automatische Excel-export: totaallijst, werklijst per koerier en pakketfiche.

De app roept `ververs_na_scan()` aan na elke in- of uitscan, zodat de
Excel-bestanden altijd overeenkomen met wat er in de database staat.
"""
from __future__ import annotations

import re
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config, db

KOP_VULLING = PatternFill("solid", fgColor="1F3864")
KOP_FONT = Font(bold=True, color="FFFFFF", size=11)
TITEL_FONT = Font(bold=True, size=14, color="1F3864")
GROEN = PatternFill("solid", fgColor="C6EFCE")
GROEN_FONT = Font(bold=True, color="1B5E20")
ORANJE = PatternFill("solid", fgColor="FFE39B")
ORANJE_FONT = Font(bold=True, color="7F4F00")
RAND = Border(*[Side(style="thin", color="BFBFBF")] * 4)

KOLOMMEN = [
    ("Barcode", "code", 22),
    ("Koerier", "koerier_naam", 18),
    ("Status", "_status", 30),
    ("Ontvanger", "ontvanger", 24),
    ("Adres", "adres", 30),
    ("Postcode", "postcode", 11),
    ("Plaats", "plaats", 18),
    ("Telefoon", "telefoon", 16),
    ("Plank/locatie", "plank", 14),
    ("Ingescand op", "in_op", 19),
    ("Ingescand door", "in_door", 16),
    ("Afgehaald op", "uit_op", 19),
    ("Uitgescand door", "uit_door", 16),
    ("Opmerking", "opmerking", 30),
]


def _statustekst(pakket: sqlite3.Row | dict) -> str:
    if pakket["status"] == config.STATUS_UIT:
        return f"AFGEHAALD ✔  {pakket['uit_op'] or ''}".strip()
    return "OPEN — nog af te halen"


def veilige_naam(tekst: str) -> str:
    """Bestandsnaam zonder tekens die op Windows/Linux problemen geven."""
    schoon = re.sub(r"[^A-Za-z0-9._-]+", "_", (tekst or "").strip())
    return schoon.strip("_") or "onbekend"


def _kopregel(ws, rij: int) -> None:
    for kolom, (titel, _, breedte) in enumerate(KOLOMMEN, start=1):
        cel = ws.cell(row=rij, column=kolom, value=titel)
        cel.fill = KOP_VULLING
        cel.font = KOP_FONT
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(kolom)].width = breedte
    ws.row_dimensions[rij].height = 24


def _datarij(ws, rij: int, pakket: sqlite3.Row | dict) -> None:
    for kolom, (_, veld, _) in enumerate(KOLOMMEN, start=1):
        waarde = _statustekst(pakket) if veld == "_status" else (pakket[veld] or "")
        cel = ws.cell(row=rij, column=kolom, value=waarde)
        cel.alignment = Alignment(vertical="center", wrap_text=(veld in ("adres", "opmerking")))
        cel.border = RAND
        if veld == "_status":
            afgehaald = pakket["status"] == config.STATUS_UIT
            cel.fill = GROEN if afgehaald else ORANJE
            cel.font = GROEN_FONT if afgehaald else ORANJE_FONT


def _schrijf_blad(ws, pakketten, titel: str, subtitel: str) -> None:
    ws.cell(row=1, column=1, value=titel).font = TITEL_FONT
    ws.cell(row=2, column=1, value=subtitel).font = Font(size=10, color="595959")
    _kopregel(ws, 4)
    for i, pakket in enumerate(pakketten, start=5):
        _datarij(ws, i, pakket)
    laatste = max(5, 4 + len(pakketten))
    ws.auto_filter.ref = f"A4:{get_column_letter(len(KOLOMMEN))}{laatste}"
    ws.freeze_panes = "A5"


def _tijdstempel() -> str:
    return datetime.now().strftime("%d-%m-%Y %H:%M")


# --- exports ----------------------------------------------------------------

def schrijf_master(conn: sqlite3.Connection) -> Path:
    """Eén Excel met alle gegevens van alle pakketten + een blad per koerier."""
    config.zorg_voor_mappen()
    alle = db.pakketten(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "Alle pakketten"
    open_ = sum(1 for p in alle if p["status"] == config.STATUS_IN)
    _schrijf_blad(
        ws, alle, "Lux 2.0 — alle pakketten",
        f"Bijgewerkt: {_tijdstempel()}  |  Totaal: {len(alle)}  |  "
        f"In depot: {open_}  |  Afgehaald: {len(alle) - open_}",
    )
    for k in db.koeriers(conn):
        rijen = [p for p in alle if p["koerier_id"] == k["id"]]
        blad = wb.create_sheet(veilige_naam(k["naam"])[:31])
        k_open = sum(1 for p in rijen if p["status"] == config.STATUS_IN)
        _schrijf_blad(
            blad, rijen, f"Werklijst {k['naam']}",
            f"Bijgewerkt: {_tijdstempel()}  |  Totaal: {len(rijen)}  |  "
            f"Nog af te halen: {k_open}  |  Afgehaald: {len(rijen) - k_open}",
        )
    pad = config.EXPORT_DIR / config.MASTER_BESTANDSNAAM
    wb.save(pad)
    return pad


def koerier_bestandsnaam(koerier_row: sqlite3.Row | dict) -> str:
    return f"{koerier_row['id']:02d}_{veilige_naam(koerier_row['naam'])}_werklijst.xlsx"


def schrijf_koerierlijst(conn: sqlite3.Connection, koerier_id: int) -> Path:
    """Aparte werklijst per koerier: open pakketten bovenaan, afgehaalde eronder."""
    config.zorg_voor_mappen()
    k = db.koerier(conn, koerier_id)
    if k is None:
        raise db.ScanFout(f"Koerier {koerier_id} bestaat niet.")
    rijen = db.pakketten(conn, koerier_id=koerier_id)
    rijen.sort(key=lambda p: (p["status"] == config.STATUS_UIT, p["in_op"] or ""))
    open_ = sum(1 for p in rijen if p["status"] == config.STATUS_IN)

    wb = Workbook()
    ws = wb.active
    ws.title = "Werklijst"
    _schrijf_blad(
        ws, rijen, f"Werklijst — {k['naam']} (Lux 2.0)",
        f"Bijgewerkt: {_tijdstempel()}  |  Totaal: {len(rijen)}  |  "
        f"Nog af te halen: {open_}  |  Afgehaald: {len(rijen) - open_}",
    )
    pad = config.KOERIER_EXPORT_DIR / koerier_bestandsnaam(k)
    wb.save(pad)
    return pad


def schrijf_alle_koerierlijsten(conn: sqlite3.Connection) -> list[Path]:
    return [schrijf_koerierlijst(conn, k["id"]) for k in db.koeriers(conn)]


def schrijf_pakket(conn: sqlite3.Connection, code: str) -> Path:
    """Pakketfiche: alle gegevens van één pakket plus de volledige scanhistorie."""
    config.zorg_voor_mappen()
    pakket = db.pakket_via_code(conn, code)
    if pakket is None:
        raise db.ScanFout(f"Pakket {code} is onbekend.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Pakket"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 46
    ws.cell(row=1, column=1, value=f"Pakket {pakket['code']}").font = TITEL_FONT
    ws.cell(row=2, column=1, value=f"Aangemaakt: {_tijdstempel()}").font = Font(
        size=10, color="595959"
    )
    rij = 4
    for titel, veld in [(t, v) for t, v, _ in KOLOMMEN]:
        label = ws.cell(row=rij, column=1, value=titel)
        label.fill = KOP_VULLING
        label.font = KOP_FONT
        waarde = _statustekst(pakket) if veld == "_status" else (pakket[veld] or "")
        cel = ws.cell(row=rij, column=2, value=waarde)
        cel.border = RAND
        if veld == "_status":
            afgehaald = pakket["status"] == config.STATUS_UIT
            cel.fill = GROEN if afgehaald else ORANJE
            cel.font = GROEN_FONT if afgehaald else ORANJE_FONT
        rij += 1

    historie = wb.create_sheet("Scanhistorie")
    for kolom, (titel, breedte) in enumerate(
        [("Tijdstip", 19), ("Actie", 10), ("Koerier", 18), ("Gescand door", 18), ("Notitie", 30)],
        start=1,
    ):
        cel = historie.cell(row=1, column=kolom, value=titel)
        cel.fill = KOP_VULLING
        cel.font = KOP_FONT
        historie.column_dimensions[get_column_letter(kolom)].width = breedte
    for i, scan in enumerate(db.scans_van_pakket(conn, pakket["id"]), start=2):
        actie = "INGESCAND" if scan["actie"] == "IN" else "AFGEHAALD"
        for kolom, waarde in enumerate(
            [scan["tijdstip"], actie, scan["koerier_naam"] or "",
             scan["gebruiker"], scan["notitie"]], start=1
        ):
            historie.cell(row=i, column=kolom, value=waarde).border = RAND
    historie.freeze_panes = "A2"

    pad = config.PAKKET_EXPORT_DIR / f"pakket_{veilige_naam(pakket['code'])}.xlsx"
    wb.save(pad)
    return pad


def ververs_na_scan(
    conn: sqlite3.Connection, koerier_id: int | None, code: str | None = None
) -> dict[str, str]:
    """Werkt na elke scan de betrokken Excel-bestanden bij."""
    paden = {"master": str(schrijf_master(conn))}
    if koerier_id is not None:
        paden["koerier"] = str(schrijf_koerierlijst(conn, koerier_id))
    if code:
        paden["pakket"] = str(schrijf_pakket(conn, code))
    return paden
